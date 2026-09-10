"""
SIH26117 — Phase 4 Multimodal Document Ingestion Pipeline Tests
Comprehensive offline test suite covering:
1. SHA-256 correctness and streaming chunking
2. Path traversal sanitization and security enforcement
3. File type & magic byte signature detection
4. Storage persistence (raw & processed) and duplicate detection
5. Digital PDF parsing, text extraction, and page provenance
6. PDF table extraction and structure normalization
7. Scanned PDF detection (NEEDS_OCR) and drawing/schematic heuristic
8. DOCX paragraph and table parsing
9. XLSX spreadsheet parsing and formula safety
10. CSV parsing, dialect sniffing, and formula injection sanitization
11. Image (PNG/JPEG) geometry validation without OCR inference
12. API endpoints (/api/v1/files/upload, /api/v1/files/{sha256}/status)
13. Industrial-like NDT sample document verification (C-101)
14. Air-gap compliance (zero external network calls)
"""

import hashlib
import io
import socket
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from PIL import Image
import docx
import openpyxl
import pypdf

try:
    from app.core.config import get_settings
    from app.main import app
    from app.services.ingestion.exceptions import (
        FileTooLargeError,
        InvalidDocumentError,
        SecurityViolationError,
        UnsupportedFileTypeError,
    )
    from app.services.ingestion.file_detector import detect_file_type
    from app.services.ingestion.hasher import (
        compute_sha256_bytes,
        compute_sha256_file,
        compute_sha256_stream,
    )
    from app.services.ingestion.models import (
        DocumentIngestionResult,
        ExtractionStatus,
        NormalizedDocument,
    )
    from app.services.ingestion.service import IngestionService
    from app.services.ingestion.storage import StorageManager, sanitize_filename
except ImportError:
    from backend.app.core.config import get_settings
    from backend.app.main import app
    from backend.app.services.ingestion.exceptions import (
        FileTooLargeError,
        InvalidDocumentError,
        SecurityViolationError,
        UnsupportedFileTypeError,
    )
    from backend.app.services.ingestion.file_detector import detect_file_type
    from backend.app.services.ingestion.hasher import (
        compute_sha256_bytes,
        compute_sha256_file,
        compute_sha256_stream,
    )
    from backend.app.services.ingestion.models import (
        DocumentIngestionResult,
        ExtractionStatus,
        NormalizedDocument,
    )
    from backend.app.services.ingestion.service import IngestionService
    from backend.app.services.ingestion.storage import StorageManager, sanitize_filename

client = TestClient(app)


@pytest.fixture(autouse=True)
def isolate_test_storage(tmp_path: Path, monkeypatch):
    """Ensure every test runs against a clean, isolated temporary storage directory."""
    test_raw = tmp_path / "raw"
    test_processed = tmp_path / "processed"
    test_raw.mkdir(parents=True, exist_ok=True)
    test_processed.mkdir(parents=True, exist_ok=True)

    settings = get_settings()
    monkeypatch.setattr(settings, "raw_data_dir", test_raw)
    monkeypatch.setattr(settings, "processed_data_dir", test_processed)


# ── Synthetic In-Memory Document Builders ─────────────────────────────────────

def build_minimal_pdf(text_lines: list[str]) -> bytes:
    """Construct an uncompressed valid PDF with digital text."""
    content_stream = "BT /F1 12 Tf 50 750 Td 14 TL "
    for line in text_lines:
        content_stream += f"({line}) ' "
    content_stream += "ET"
    content_bytes = content_stream.encode("latin1")
    stream_len = len(content_bytes)

    pdf = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>
endobj
4 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
5 0 obj
<< /Length {stream_len} >>
stream
{content_stream}
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000234 00000 n 
0000000305 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
400
%%EOF
"""
    return pdf.encode("latin1")


def build_table_pdf() -> bytes:
    """Construct a valid PDF containing grid lines and table cells."""
    content = """BT /F1 10 Tf
60 710 Td (Point) Tj
160 0 Td (Reading) Tj
-160 -25 Td (P-01) Tj
160 0 Td (11.2) Tj
-160 -25 Td (P-02) Tj
160 0 Td (10.9) Tj
ET
50 725 m 250 725 l S
50 700 m 250 700 l S
50 675 m 250 675 l S
50 650 m 250 650 l S
50 725 m 50 650 l S
150 725 m 150 650 l S
250 725 m 250 650 l S
"""
    stream_bytes = content.encode("latin1")
    pdf = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>
endobj
4 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
5 0 obj
<< /Length {len(stream_bytes)} >>
stream
{content}endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000234 00000 n 
0000000305 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
400
%%EOF
"""
    return pdf.encode("latin1")


def build_blank_pdf() -> bytes:
    """Construct a blank PDF without any text layer (simulating a scan)."""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def build_drawing_pdf() -> bytes:
    """Construct a wide landscape page with dense vector lines simulating a P&ID schematic."""
    content = ""
    # Generate 40 vector lines
    for i in range(40):
        y = 100 + i * 15
        content += f"50 {y} m 950 {y} l S\n"
    stream_bytes = content.encode("latin1")
    pdf = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 1000 600] /Resources << >> /Contents 4 0 R >>
endobj
4 0 obj
<< /Length {len(stream_bytes)} >>
stream
{content}endstream
endobj
xref
0 5
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000210 00000 n 
trailer
<< /Size 5 /Root 1 0 R >>
startxref
350
%%EOF
"""
    return pdf.encode("latin1")


def build_docx_sample() -> bytes:
    """Construct a valid in-memory DOCX document."""
    doc = docx.Document()
    doc.add_heading("MRPL Technical Report", level=1)
    doc.add_paragraph("First paragraph describing refinery piping inspection.")
    doc.add_heading("Section 2: Ultrasonic Readings", level=2)
    doc.add_paragraph("Measurements collected at CDU-101.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Point"
    table.cell(0, 1).text = "Thickness_mm"
    table.cell(1, 0).text = "P-101"
    table.cell(1, 1).text = "11.2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_xlsx_sample() -> bytes:
    """Construct a valid in-memory XLSX workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NDT_Readings"
    ws.append(["Inspection_Point", "Nominal_mm", "Actual_mm"])
    ws.append(["Point_A", 12.0, 11.2])
    ws.append(["Point_B", 12.0, 10.9])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_image_sample(fmt: str = "PNG") -> bytes:
    """Construct a valid in-memory image."""
    img = Image.new("RGB", (120, 80), color=(30, 80, 150))
    buf = io.BytesIO()
    img.save(buf, format=fmt)
    return buf.getvalue()


# ── Tests: SHA-256 Hashing ───────────────────────────────────────────────────

def test_sha256_correctness():
    """Verify SHA-256 computation matches standard hashlib reference."""
    sample = b"MRPL Sovereign Ingestion Test Payload"
    expected = hashlib.sha256(sample).hexdigest()
    assert compute_sha256_bytes(sample) == expected


def test_sha256_stream_chunking():
    """Verify streaming SHA-256 reader chunks properly without altering position."""
    large_data = b"Industrial Refinery Log Entry\n" * 5000
    expected = hashlib.sha256(large_data).hexdigest()

    stream = io.BytesIO(large_data)
    digest = compute_sha256_stream(stream, chunk_size=512)
    assert digest == expected
    assert stream.tell() == 0  # Stream seek position preserved


def test_sha256_file(tmp_path: Path):
    """Verify file-based SHA-256 computation."""
    f = tmp_path / "test_file.bin"
    f.write_bytes(b"File on disk hashing verification")
    expected = hashlib.sha256(b"File on disk hashing verification").hexdigest()
    assert compute_sha256_file(f) == expected


# ── Tests: Filename Sanitization & Security ───────────────────────────────────

def test_filename_sanitization_normal():
    """Verify clean filenames are preserved."""
    assert sanitize_filename("report_2026.pdf") == "report_2026.pdf"


def test_path_traversal_relative():
    """Verify relative traversal characters are stripped."""
    assert sanitize_filename("../../secret.pdf") == "secret.pdf"
    assert sanitize_filename("..\\..\\passwords.txt") == "passwords.txt"


def test_path_traversal_absolute_unix():
    """Verify Unix absolute path elements are reduced to leaf basename."""
    assert sanitize_filename("/etc/shadow.pdf") == "shadow.pdf"


def test_path_traversal_absolute_windows():
    """Verify Windows absolute paths are reduced to leaf basename."""
    assert sanitize_filename("C:\\Windows\\system32\\cmd.pdf") == "cmd.pdf"


def test_null_byte_rejection():
    """Verify null bytes trigger a SecurityViolationError."""
    with pytest.raises(SecurityViolationError):
        sanitize_filename("exploit\x00.pdf")


def test_empty_filename_fallback():
    """Verify empty or space-only filenames fall back to default."""
    assert sanitize_filename("") == "unnamed_artifact"
    assert sanitize_filename("   ") == "unnamed_artifact"
    assert sanitize_filename("../../") == "unnamed_artifact"


# ── Tests: File Type Detection & Signature Verification ───────────────────────

def test_detect_valid_pdf():
    """Verify valid PDF header signature is detected."""
    pdf_bytes = build_minimal_pdf(["Sample Text"])
    media_type, ext = detect_file_type(pdf_bytes, "doc.pdf")
    assert media_type == "application/pdf"
    assert ext == "pdf"


def test_detect_pdf_missing_header():
    """Verify PDF with invalid magic header is rejected."""
    with pytest.raises(InvalidDocumentError):
        detect_file_type(b"NOT_A_PDF_STREAM", "doc.pdf")


def test_detect_png():
    """Verify valid PNG magic bytes are detected."""
    png_bytes = build_image_sample("PNG")
    media_type, ext = detect_file_type(png_bytes, "scan.png")
    assert media_type == "image/png"
    assert ext == "png"


def test_detect_jpeg():
    """Verify valid JPEG magic bytes are detected."""
    jpeg_bytes = build_image_sample("JPEG")
    media_type, ext = detect_file_type(jpeg_bytes, "photo.jpg")
    assert media_type == "image/jpeg"
    assert ext == "jpg"


def test_detect_unsupported_extension():
    """Verify unsupported file types raise UnsupportedFileTypeError."""
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(b"some text", "script.sh")


def test_detect_dangerous_executable():
    """Verify dangerous executables are rejected immediately."""
    with pytest.raises(UnsupportedFileTypeError):
        detect_file_type(b"MZ\x90\x00", "installer.exe")


# ── Tests: Storage Manager & Duplicate Handling ───────────────────────────────

def test_raw_persistence_and_duplicate(tmp_path: Path):
    """Verify raw persistence and identical duplicate detection."""
    settings = get_settings()
    storage = StorageManager(settings)

    payload = b"Unique Raw Content for Ingestion"
    sha256 = compute_sha256_bytes(payload)

    # First write
    path1, is_dup1 = storage.store_raw_file(sha256, "sample.pdf", payload)
    assert path1.exists()
    assert not is_dup1
    assert path1.read_bytes() == payload

    # Second write (duplicate)
    path2, is_dup2 = storage.store_raw_file(sha256, "sample.pdf", payload)
    assert path2 == path1
    assert is_dup2 is True


# ── Tests: Service Ingestion Workflows ────────────────────────────────────────

def test_ingest_digital_pdf():
    """Verify text and page count extraction from a digital PDF."""
    service = IngestionService()
    pdf_bytes = build_minimal_pdf(["Equipment: C-101", "Measured Thickness: 11.2 mm"])

    result = service.ingest_file(pdf_bytes, "inspection_c101.pdf")
    assert result.status == ExtractionStatus.SUCCESS
    assert result.page_count == 1
    assert result.is_scanned is False
    assert result.needs_ocr is False
    assert result.media_type == "application/pdf"
    assert "inspection_c101" in result.filename

    doc = result.normalized_document
    assert doc is not None
    assert len(doc.pages) == 1
    assert "C-101" in doc.pages[0].text
    assert doc.pages[0].provenance.source_filename == "inspection_c101.pdf"


def test_ingest_pdf_with_table():
    """Verify table extraction from digital PDF with cell structures."""
    service = IngestionService()
    pdf_bytes = build_table_pdf()

    result = service.ingest_file(pdf_bytes, "ndt_table.pdf")
    assert result.status == ExtractionStatus.SUCCESS
    assert result.table_count >= 1
    table = result.normalized_document.tables[0]
    assert table.row_count >= 2
    assert table.extraction_method == "pdfplumber"
    assert table.provenance.page_number == 1


def test_ingest_scanned_pdf_needs_ocr():
    """Verify scanned blank PDF is detected and marked NEEDS_OCR without calling OCR."""
    service = IngestionService()
    blank_pdf = build_blank_pdf()

    result = service.ingest_file(blank_pdf, "scanned_ndt_log.pdf")
    assert result.is_scanned is True
    assert result.needs_ocr is True
    assert result.status == ExtractionStatus.NEEDS_OCR
    assert any("OCR" in w for w in result.warnings)


def test_ingest_drawing_heuristic():
    """Verify drawing schematic heuristic flags wide page with vector lines."""
    service = IngestionService()
    drawing_pdf = build_drawing_pdf()

    result = service.ingest_file(drawing_pdf, "pid_schematic.pdf")
    assert result.is_drawing is True
    assert any("schematic" in w or "drawing" in w for w in result.warnings)


def test_ingest_docx():
    """Verify DOCX paragraph, heading, and table extraction."""
    service = IngestionService()
    docx_bytes = build_docx_sample()

    result = service.ingest_file(docx_bytes, "technical_memo.docx")
    assert result.status == ExtractionStatus.SUCCESS
    assert result.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    doc = result.normalized_document
    assert len(doc.sections) >= 2
    assert any(s.heading == "MRPL Technical Report" for s in doc.sections)
    assert len(doc.tables) == 1
    assert doc.tables[0].headers == ["Point", "Thickness_mm"]


def test_ingest_xlsx():
    """Verify XLSX spreadsheet sheet and table extraction."""
    service = IngestionService()
    xlsx_bytes = build_xlsx_sample()

    result = service.ingest_file(xlsx_bytes, "readings.xlsx")
    assert result.status == ExtractionStatus.SUCCESS
    doc = result.normalized_document
    assert len(doc.sheets) == 1
    assert doc.sheets[0].sheet_name == "NDT_Readings"
    assert len(doc.tables) == 1
    assert doc.tables[0].headers == ["Inspection_Point", "Nominal_mm", "Actual_mm"]


def test_ingest_csv():
    """Verify CSV parsing with delimiter sniffing and row provenance."""
    service = IngestionService()
    csv_bytes = b"Tag,Reading_mm,Status\nC-101,11.2,Normal\nC-102,10.9,Normal\n"

    result = service.ingest_file(csv_bytes, "piping_data.csv")
    assert result.status == ExtractionStatus.SUCCESS
    assert result.table_count == 1
    table = result.normalized_document.tables[0]
    assert table.headers == ["Tag", "Reading_mm", "Status"]
    assert table.row_count == 2
    assert table.provenance.sheet_name == "CSV_Data"


def test_csv_formula_injection_sanitization():
    """Verify CSV formula injection characters are neutralized into string literals."""
    service = IngestionService()
    evil_csv = b"Command,Value\n=cmd|' /C calc'!A0,123\n+SUM(A1:A10),456\n"

    result = service.ingest_file(evil_csv, "suspicious.csv")
    table = result.normalized_document.tables[0]
    first_cell = table.rows[0][0]
    # Sanitized cell must start with single quote to prevent active spreadsheet execution
    assert first_cell.startswith("'=")


def test_ingest_image_png():
    """Verify PNG image properties extracted and flagged for Phase 5 OCR."""
    service = IngestionService()
    png_bytes = build_image_sample("PNG")

    result = service.ingest_file(png_bytes, "inspection_photo.png")
    assert result.status == ExtractionStatus.SUCCESS
    assert result.needs_ocr is True
    doc = result.normalized_document
    assert doc.image_info is not None
    assert doc.image_info.width == 120
    assert doc.image_info.height == 80
    assert doc.image_info.format == "PNG"


def test_file_too_large_rejection():
    """Verify file exceeding maximum upload size is rejected."""
    service = IngestionService()
    # Temporarily set max size to 100 bytes
    service.settings.max_upload_size_bytes = 100
    try:
        with pytest.raises(FileTooLargeError):
            service.ingest_file(b"X" * 200, "large.pdf")
    finally:
        # Restore default 50 MB
        service.settings.max_upload_size_bytes = 52428800


def test_empty_file_rejection():
    """Verify 0-byte file is rejected as INVALID_DOCUMENT."""
    service = IngestionService()
    with pytest.raises(InvalidDocumentError):
        service.ingest_file(b"", "empty.pdf")


# ── Tests: Synthetic Industrial Sample NDT Document ───────────────────────────

def test_synthetic_industrial_ndt_sample():
    """
    Ingest data/samples/corrosion_inspection_c101.pdf and verify:
    - Equipment identifier: C-101
    - Inspection Date: 2026-03-15
    - Nominal Thickness: 12.0 mm
    - Ultrasonic readings: 11.2, 10.9, 10.5, 10.1
    - Table columns and 4 rows
    - Page provenance correctly attached
    """
    sample_path = Path(__file__).resolve().parent.parent.parent / "data" / "samples" / "corrosion_inspection_c101.pdf"
    assert sample_path.is_file(), f"Sample file missing: {sample_path}"

    pdf_bytes = sample_path.read_bytes()
    service = IngestionService()
    result = service.ingest_file(pdf_bytes, "corrosion_inspection_c101.pdf")

    assert result.status == ExtractionStatus.SUCCESS
    assert result.page_count == 1
    assert result.table_count >= 1
    assert result.is_scanned is False

    doc = result.normalized_document
    assert doc is not None
    page_text = doc.pages[0].text

    assert "C-101" in page_text
    assert "2026-03-15" in page_text
    assert "12.0 mm" in page_text

    table = doc.tables[0]
    assert table.row_count == 4
    # Verify readings exist in rows
    cell_values = [str(cell) for row in table.rows for cell in row]
    assert "11.2" in cell_values
    assert "10.9" in cell_values
    assert "10.5" in cell_values
    assert "10.1" in cell_values
    assert table.provenance.page_number == 1


# ── Tests: API Upload Endpoints ───────────────────────────────────────────────

def test_api_upload_pdf_success():
    """Verify POST /api/v1/files/upload with PDF returns 201 Created."""
    pdf_bytes = build_minimal_pdf(["API Gateway PDF Ingestion Test"])
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("api_test.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 201
    data = response.json()
    assert "document_id" in data
    assert "sha256" in data
    assert data["filename"] == "api_test.pdf"
    assert data["page_count"] == 1
    assert data["is_duplicate"] is False


def test_api_upload_duplicate():
    """Verify second identical upload returns is_duplicate = True."""
    pdf_bytes = build_minimal_pdf(["Duplicate Detection Payload"])
    # First upload
    res1 = client.post(
        "/api/v1/files/upload",
        files={"file": ("first.pdf", pdf_bytes, "application/pdf")},
    )
    assert res1.status_code == 201
    assert res1.json()["is_duplicate"] is False

    # Second upload with same bytes
    res2 = client.post(
        "/api/v1/files/upload",
        files={"file": ("first.pdf", pdf_bytes, "application/pdf")},
    )
    assert res2.status_code == 201
    assert res2.json()["is_duplicate"] is True
    assert res2.json()["sha256"] == res1.json()["sha256"]


def test_api_upload_path_traversal_sanitization():
    """Verify path traversal filenames in multipart uploads are safely sanitized."""
    pdf_bytes = build_minimal_pdf(["Path Traversal Test"])
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("../../evil_traversal.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "evil_traversal.pdf"
    assert ".." not in data["raw_path"]


def test_api_upload_unsupported_extension():
    """Verify unsupported file extension returns 415 with UNSUPPORTED_FILE_TYPE."""
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("malicious.exe", b"MZ_FAKE_EXE_BINARY", "application/octet-stream")},
    )
    assert response.status_code == 415
    data = response.json()
    assert data["error"]["code"] == "UNSUPPORTED_FILE_TYPE"


def test_api_upload_corrupted_pdf():
    """Verify corrupt PDF content returns 422 with INVALID_DOCUMENT."""
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("corrupt.pdf", b"%PDF-corrupted-binary-trash", "application/pdf")},
    )
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == "INVALID_DOCUMENT"


def test_api_get_document_status():
    """Verify GET /api/v1/files/{sha256}/status returns stored document representation."""
    pdf_bytes = build_minimal_pdf(["Status Query Payload"])
    upload_res = client.post(
        "/api/v1/files/upload",
        files={"file": ("status_test.pdf", pdf_bytes, "application/pdf")},
    )
    sha256 = upload_res.json()["sha256"]

    status_res = client.get(f"/api/v1/files/{sha256}/status")
    assert status_res.status_code == 200
    doc_data = status_res.json()
    assert doc_data["sha256"] == sha256
    assert doc_data["sanitized_filename"] == "status_test.pdf"


def test_api_get_document_status_not_found():
    """Verify querying an unknown SHA-256 returns 404."""
    response = client.get("/api/v1/files/0000000000000000000000000000000000000000000000000000000000000000/status")
    assert response.status_code == 404


# ── Air-Gap & Sovereignty Verification ────────────────────────────────────────

def test_air_gap_no_outbound_network_calls(monkeypatch):
    """
    Verify that document ingestion makes zero network calls.
    Patches socket.create_connection to fail if any external connection is attempted.
    """
    def forbidden_connect(*args, **kwargs):
        raise AssertionError("Air-gap violation: Attempted outbound socket connection during ingestion!")

    monkeypatch.setattr(socket, "create_connection", forbidden_connect)

    service = IngestionService()
    pdf_bytes = build_minimal_pdf(["Air-Gap Ingestion Compliance"])
    result = service.ingest_file(pdf_bytes, "air_gap_doc.pdf")
    assert result.status == ExtractionStatus.SUCCESS


def test_corrupt_docx_error():
    """Verify invalid DOCX zip container raises InvalidDocumentError."""
    service = IngestionService()
    # Malformed zip header
    with pytest.raises(InvalidDocumentError):
        service.ingest_file(b"PK\x03\x04corrupted_zip_bytes", "bad.docx")


def test_corrupt_xlsx_error():
    """Verify invalid XLSX zip container raises InvalidDocumentError."""
    service = IngestionService()
    with pytest.raises(InvalidDocumentError):
        service.ingest_file(b"PK\x03\x04corrupted_zip_bytes", "bad.xlsx")


def test_sha256_file_not_found():
    """Verify compute_sha256_file raises FileNotFoundError for missing path."""
    with pytest.raises(FileNotFoundError):
        compute_sha256_file("nonexistent_path_to_artifact.pdf")


def test_api_upload_docx_success():
    """Verify POST /api/v1/files/upload with DOCX returns 201 Created."""
    docx_bytes = build_docx_sample()
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("memo.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "memo.docx"
    assert data["status"] == "success"
    assert data["page_count"] >= 2


def test_api_upload_xlsx_success():
    """Verify POST /api/v1/files/upload with XLSX returns 201 Created."""
    xlsx_bytes = build_xlsx_sample()
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("sheet.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "sheet.xlsx"
    assert data["table_count"] == 1


def test_api_upload_csv_success():
    """Verify POST /api/v1/files/upload with CSV returns 201 Created."""
    csv_bytes = b"Col1,Col2\nVal1,Val2\n"
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("data.csv", csv_bytes, "text/csv")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "data.csv"
    assert data["table_count"] == 1


def test_api_upload_image_png_success():
    """Verify POST /api/v1/files/upload with PNG returns 201 Created and needs_ocr flag."""
    png_bytes = build_image_sample("PNG")
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("chart.png", png_bytes, "image/png")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "chart.png"
    assert data["needs_ocr"] is True


def test_api_upload_image_jpg_success():
    """Verify POST /api/v1/files/upload with JPG returns 201 Created."""
    jpg_bytes = build_image_sample("JPEG")
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("photo.jpg", jpg_bytes, "image/jpeg")},
    )
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "photo.jpg"
    assert data["needs_ocr"] is True


def test_api_upload_oversized_file(monkeypatch):
    """Verify API returns 413 when file exceeds configured limit."""
    settings = get_settings()
    monkeypatch.setattr(settings, "max_upload_size_bytes", 50)
    pdf_bytes = build_minimal_pdf(["This string will easily exceed fifty bytes limit"])
    response = client.post(
        "/api/v1/files/upload",
        files={"file": ("oversized.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 413
    data = response.json()
    assert data["error"]["code"] == "FILE_TOO_LARGE"
