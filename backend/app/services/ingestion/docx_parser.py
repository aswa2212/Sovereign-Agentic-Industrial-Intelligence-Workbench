"""
SIH26117 — DOCX Parser Module
Extracts structured paragraphs, headings, and tables from Word documents.
Preserves linear section order and paragraph-level provenance.
"""

import io
from typing import Any, Dict, List, Tuple
import docx

try:
    from app.services.ingestion.exceptions import InvalidDocumentError
    from app.services.ingestion.models import (
        DocumentProvenance,
        DocumentSection,
        ParsedTable,
    )
    from app.services.ingestion.table_extractor import clean_cell_value
except ImportError:
    from backend.app.services.ingestion.exceptions import InvalidDocumentError
    from backend.app.services.ingestion.models import (
        DocumentProvenance,
        DocumentSection,
        ParsedTable,
    )
    from backend.app.services.ingestion.table_extractor import clean_cell_value


class DOCXParser:
    """Parses DOCX documents into structured sections, headings, and tables."""

    def parse(
        self,
        content: bytes,
        sha256: str,
        filename: str,
    ) -> Tuple[List[DocumentSection], List[ParsedTable], Dict[str, Any]]:
        """
        Parse raw DOCX bytes.

        Returns:
            Tuple: (sections, tables, metadata)
        """
        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise InvalidDocumentError(
                f"Failed to parse DOCX document: {str(e)}",
                details={"filename": filename},
            )

        # 1. Extract core document metadata
        metadata: Dict[str, Any] = {}
        try:
            core_props = doc.core_properties
            if core_props:
                if core_props.title:
                    metadata["title"] = core_props.title
                if core_props.author:
                    metadata["author"] = core_props.author
                if core_props.created:
                    metadata["created"] = str(core_props.created)
                if core_props.modified:
                    metadata["modified"] = str(core_props.modified)
                if core_props.revision:
                    metadata["revision"] = core_props.revision
        except Exception:
            pass

        # 2. Extract paragraphs and headings
        sections: List[DocumentSection] = []
        for idx, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else "Normal"
            heading_level = None
            if style_name.startswith("Heading"):
                try:
                    heading_level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    heading_level = 1

            provenance = DocumentProvenance(
                source_filename=filename,
                source_sha256=sha256,
                section_index=idx,
            )

            section = DocumentSection(
                section_index=idx,
                heading_level=heading_level,
                heading=text if heading_level else None,
                text=text,
                tables=[],
                provenance=provenance,
            )
            sections.append(section)

        # 3. Extract tables
        tables: List[ParsedTable] = []
        for t_idx, table in enumerate(doc.tables):
            raw_rows = []
            for row in table.rows:
                row_cells = [clean_cell_value(cell.text) for cell in row.cells]
                raw_rows.append(row_cells)

            if not raw_rows:
                continue

            # First row as header
            raw_headers = raw_rows[0]
            headers = [
                str(h).strip() if str(h).strip() else f"Col_{c_idx + 1}"
                for c_idx, h in enumerate(raw_headers)
            ]
            data_rows = raw_rows[1:] if len(raw_rows) > 1 else []

            table_id = f"table_{sha256[:8]}_docx_{t_idx + 1}"
            provenance = DocumentProvenance(
                source_filename=filename,
                source_sha256=sha256,
                section_index=t_idx,
            )

            parsed_table = ParsedTable(
                table_id=table_id,
                headers=headers,
                rows=data_rows,
                row_count=len(data_rows),
                col_count=len(headers),
                extraction_method="python-docx",
                confidence=1.0,
                provenance=provenance,
            )
            tables.append(parsed_table)

        metadata["paragraph_count"] = len(doc.paragraphs)
        metadata["table_count"] = len(tables)

        return sections, tables, metadata
